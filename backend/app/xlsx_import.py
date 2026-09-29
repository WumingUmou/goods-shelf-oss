"""Parse the 谷子 Excel template (header row + embedded images) into plain rows."""
import posixpath
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from xml.etree import ElementTree as ET

import openpyxl

from .models import DEFAULT_STATUS, STATUSES
from .seed import normalize_name

NS = {
    "m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
    "rd": "http://schemas.microsoft.com/office/spreadsheetml/2017/richdata",
    "rv": "http://schemas.microsoft.com/office/spreadsheetml/2022/richvaluerel",
    "rel": "http://schemas.openxmlformats.org/package/2006/relationships",
    "xlrd": "http://schemas.microsoft.com/office/spreadsheetml/2017/richdata",
}
R_ID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"

# header text → field; first match wins, so list aliases
COLUMNS = {
    "name": ["谷子名称", "名称"],
    "image": ["图片"],
    "series": ["系列"],
    "character": ["角色 / 推", "角色"],
    "kind": ["种类"],
    "bag_size": ["属性", "自封袋尺寸", "自封袋"],
    "spec": ["尺寸 / 规格", "尺寸/规格", "规格"],
    "status": ["状态"],
    "quantity": ["数量"],
    "note": ["备注"],
}
SPLIT = re.compile(r"[、,，/\n]+")
NO_SERIES = "未分类"


@dataclass
class Row:
    row: int
    name: str
    series: str
    kinds: list[str]
    characters: list[str]
    bag_size: str | None
    spec: str | None
    status: str
    quantity: int
    note: str | None
    image: bytes | None = field(default=None, repr=False)
    warnings: list[str] = field(default_factory=list)


def _split(v) -> list[str]:
    if v is None:
        return []
    out = [normalize_name(p) for p in SPLIT.split(str(v))]
    return list(dict.fromkeys(p for p in out if p))


def _text(v) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def _cell_images(path: Path, sheet_index: int) -> dict[int, bytes]:
    """Images placed *in* cells (Excel 'Place in Cell' → richData). row → bytes."""
    out: dict[int, bytes] = {}
    with zipfile.ZipFile(path) as z:
        names = set(z.namelist())
        if "xl/metadata.xml" not in names or "xl/richData/rdrichvalue.xml" not in names:
            return out
        # sheet path for this index
        wb = ET.fromstring(z.read("xl/workbook.xml"))
        wb_rels = ET.fromstring(z.read("xl/_rels/workbook.xml.rels"))
        sheet = wb.findall("m:sheets/m:sheet", NS)[sheet_index]
        target = next(r.get("Target") for r in wb_rels if r.get("Id") == sheet.get(R_ID))
        sheet_path = posixpath.normpath(posixpath.join("xl", target.lstrip("/").removeprefix("xl/")))

        meta = ET.fromstring(z.read("xl/metadata.xml"))
        fut = [bk.find(".//xlrd:rvb", NS).get("i") for bk in meta.findall("m:futureMetadata/m:bk", NS)]
        vm_to_rv = []
        for bk in meta.findall("m:valueMetadata/m:bk", NS):
            rc = bk.find("m:rc", NS)
            vm_to_rv.append(int(fut[int(rc.get("v"))]))
        rvs = ET.fromstring(z.read("xl/richData/rdrichvalue.xml")).findall("rd:rv", NS)
        rels = ET.fromstring(z.read("xl/richData/richValueRel.xml")).findall("rv:rel", NS)
        rel_targets = {
            r.get("Id"): r.get("Target")
            for r in ET.fromstring(z.read("xl/richData/_rels/richValueRel.xml.rels"))
        }
        sheet_xml = ET.fromstring(z.read(sheet_path))
        for c in sheet_xml.iter(f"{{{NS['m']}}}c"):
            vm = c.get("vm")
            if not vm:
                continue
            try:
                rv = rvs[vm_to_rv[int(vm) - 1]]
                rel_idx = int(rv.find("rd:v", NS).text)
                tgt = rel_targets[rels[rel_idx].get(R_ID)]
                media = posixpath.normpath(posixpath.join("xl/richData", tgt))
                row = int(re.sub(r"[A-Z]+", "", c.get("r")))
                out[row] = z.read(media)
            except (IndexError, KeyError, ValueError, AttributeError):
                continue
    return out


def parse(path: Path) -> list[Row]:
    wb = openpyxl.load_workbook(path)
    ws = wb.worksheets[0]
    header = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    col: dict[str, int] = {}
    for key, aliases in COLUMNS.items():
        for a in aliases:
            if a in header:
                col[key] = header.index(a)
                break
    if "name" not in col:
        raise ValueError("第一个工作表缺少「谷子名称」列")

    floating: dict[int, bytes] = {}
    for im in ws._images:
        r = im.anchor._from.row + 1
        floating.setdefault(r, im._data())
    in_cell = _cell_images(path, 0)

    def get(values, key):
        i = col.get(key)
        return values[i] if i is not None and i < len(values) else None

    rows: list[Row] = []
    for r_idx, values in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        name = _text(get(values, "name"))
        if not name:
            continue
        warnings = []
        series = _text(get(values, "series"))
        if not series:
            series = name.split(" - ", 1)[0].strip() if " - " in name else NO_SERIES
        status = _text(get(values, "status")) or DEFAULT_STATUS
        if status not in STATUSES:
            warnings.append(f"未知状态「{status}」，按{DEFAULT_STATUS}导入")
            status = DEFAULT_STATUS
        try:
            qty = int(get(values, "quantity") or 1)
        except (TypeError, ValueError):
            qty = 1
            warnings.append("数量无法识别，按 1 导入")
        kinds = _split(get(values, "kind"))
        if not kinds:
            kinds = ["其他"]
            warnings.append("未填种类，归入「其他」")
        image = floating.get(r_idx) or in_cell.get(r_idx)
        if image is None:
            warnings.append("没有图片")
        rows.append(Row(
            row=r_idx, name=name, series=normalize_name(series), kinds=kinds,
            characters=_split(get(values, "character")),
            bag_size=_text(get(values, "bag_size")), spec=_text(get(values, "spec")),
            status=status, quantity=max(1, qty), note=_text(get(values, "note")),
            image=image, warnings=warnings,
        ))
    return rows
